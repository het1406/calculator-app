"""Generate the Word review draft from the rationale and test evidence."""
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
document = Document()
section = document.sections[0]
section.top_margin = section.bottom_margin = Inches(0.7)
style = document.styles['Normal']
style.font.name = 'Calibri'
style.font.size = Pt(11)
style.paragraph_format.space_after = Pt(8)


def markdown(path, heading_offset=0):
    lines = path.read_text().splitlines()
    paragraph = []

    def flush():
        if paragraph:
            document.add_paragraph(' '.join(paragraph))
            paragraph.clear()

    i = 0
    while i < len(lines):
        line = lines[i]
        if line.startswith('|'):
            flush()
            rows = []
            while i < len(lines) and lines[i].startswith('|'):
                cells = [c.strip() for c in lines[i].strip('|').split('|')]
                if not all(set(c) <= set('-: ') for c in cells):
                    rows.append(cells)
                i += 1
            table = document.add_table(rows=1, cols=len(rows[0]))
            table.style = 'Light Shading Accent 1'
            for c, value in zip(table.rows[0].cells, rows[0]):
                c.text = value
            for row in rows[1:]:
                for c, value in zip(table.add_row().cells, row):
                    c.text = value
            continue
        if line.startswith('#'):
            flush()
            depth = len(line) - len(line.lstrip('#'))
            document.add_heading(line.lstrip('# ').strip(), min(3, depth - 1 + heading_offset))
        elif line.startswith('- '):
            flush()
            document.add_paragraph(line[2:], style='List Bullet')
        elif not line.strip():
            flush()
        else:
            paragraph.append(line)
        i += 1
    flush()


markdown(ROOT / 'docs/implementation.md')
document.add_page_break()
markdown(ROOT / 'docs/agent-test-drive.md', heading_offset=1)
verification = ROOT / 'docs/device-verification.md'
if verification.exists():
    document.add_page_break()
    markdown(verification, heading_offset=1)
    for name in ['result.png', 'dark.png', 'error.png']:
        screenshot = ROOT / 'docs/screenshots' / name
        if screenshot.exists():
            document.add_picture(str(screenshot), width=Inches(2.5))
            document.add_paragraph(name.replace('.png', '').capitalize())
footer = section.footer.paragraphs[0]
footer.text = 'HetJani • CSC 4360 • Assignment 01'
footer.runs[0].font.size = Pt(9)
footer.runs[0].font.color.rgb = RGBColor.from_string('666666')
output = ROOT / 'submission/HetJani_Implementation.docx'
document.save(output)
print(output)
